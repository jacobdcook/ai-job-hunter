"""
Merge existing Excel files into a single combined master file.
Combines jobs_master.xlsx and pge_jobs_master.xlsx into one file with tabs.
"""

import pandas as pd
import os
from openpyxl.styles import Font, PatternFill, Alignment

def merge_excel_files():
    """Merge existing Excel files into jobs_master.xlsx with tabs."""
    
    all_jobs = []
    pge_jobs = []
    smud_jobs = []
    
    # Read jobs_master.xlsx if it exists
    if os.path.exists("jobs_master.xlsx"):
        print("📖 Reading jobs_master.xlsx...")
        try:
            df = pd.read_excel("jobs_master.xlsx", sheet_name=None)
            for sheet_name, sheet_df in df.items():
                if 'Source' in sheet_df.columns:
                    all_jobs.append(sheet_df)
                    pge_jobs.append(sheet_df[sheet_df['Source'] == 'PG&E'])
                    smud_jobs.append(sheet_df[sheet_df['Source'] == 'SMUD'])
                else:
                    # Old format without Source column
                    all_jobs.append(sheet_df)
        except Exception as e:
            print(f"  Warning: Could not read jobs_master.xlsx: {e}")
    
    # Read pge_jobs_master.xlsx if it exists
    if os.path.exists("pge_jobs_master.xlsx"):
        print("📖 Reading pge_jobs_master.xlsx...")
        try:
            df = pd.read_excel("pge_jobs_master.xlsx", sheet_name=None)
            for sheet_name, sheet_df in df.items():
                # Add Source column if missing
                if 'Source' not in sheet_df.columns:
                    sheet_df['Source'] = 'PG&E'
                all_jobs.append(sheet_df)
                pge_jobs.append(sheet_df[sheet_df['Source'] == 'PG&E'])
        except Exception as e:
            print(f"  Warning: Could not read pge_jobs_master.xlsx: {e}")
    
    if not all_jobs:
        print("❌ No Excel files found to merge.")
        return
    
    # Combine all dataframes
    print("🔗 Combining data...")
    df_all = pd.concat(all_jobs, ignore_index=True)
    
    # Deduplicate by link (keep first occurrence)
    if 'Apply Link' in df_all.columns:
        df_all = df_all.drop_duplicates(subset=['Apply Link'], keep='first')
    elif 'link' in df_all.columns:
        df_all = df_all.drop_duplicates(subset=['link'], keep='first')
    
    # Ensure Source column exists
    if 'Source' not in df_all.columns:
        df_all['Source'] = 'Unknown'
    
    # Sort by score (highest first)
    if 'Score' in df_all.columns:
        df_all = df_all.sort_values(
            by='Score',
            ascending=False,
            na_position='last'
        )
    elif 'match_score' in df_all.columns:
        df_all = df_all.sort_values(
            by='match_score',
            ascending=False,
            na_position='last'
        )
    
    # Create separate dataframes per source
    df_pge = df_all[df_all['Source'] == 'PG&E'].copy() if 'Source' in df_all.columns else pd.DataFrame()
    df_smud = df_all[df_all['Source'] == 'SMUD'].copy() if 'Source' in df_all.columns else pd.DataFrame()
    
    # Standardize column names
    column_mapping = {
        'link': 'Apply Link',
        'title': 'Job Title',
        'match_score': 'Score',
    }
    for old_col, new_col in column_mapping.items():
        if old_col in df_all.columns and new_col not in df_all.columns:
            df_all[new_col] = df_all[old_col]
    
    # Ensure we have the standard columns
    standard_cols = ['Priority', 'Score', 'Job Title', 'Location', 'Source', 'Apply Link',
                     'Missing Skills', 'AI Analysis', 'Date Posted', 'Last Seen (Active)', 'Status']
    
    # Add Priority if missing
    if 'Priority' not in df_all.columns and 'Score' in df_all.columns:
        def get_priority(score):
            if pd.isna(score) or score is None:
                return "Not Analyzed"
            elif score >= 8:
                return "HIGH PRIORITY ⭐"
            elif score >= 5:
                return "MEDIUM PRIORITY ✅"
            elif score >= 1:
                return "LOW PRIORITY 📋"
            else:
                return "SKIP"
        df_all['Priority'] = df_all['Score'].apply(get_priority)
    
    # Reorder columns to match standard format
    available_cols = [col for col in standard_cols if col in df_all.columns]
    df_all = df_all[available_cols]
    
    # Format function
    def format_worksheet(worksheet, link_col_idx=5):
        """Apply formatting to a worksheet."""
        # Auto-adjust column widths
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
        
        # Make links clickable
        link_font = Font(color="0563C1", underline="single")
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            if len(row) > link_col_idx:
                link_cell = row[link_col_idx]
                if link_cell.value and isinstance(link_cell.value, str) and link_cell.value.startswith('http'):
                    link_cell.font = link_font
                    link_cell.hyperlink = link_cell.value
        
        # Format header row
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Write to new combined file
    output_file = "jobs_master.xlsx"
    print(f"💾 Writing combined file: {output_file}...")
    
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        # Write All Jobs tab
        df_all.to_excel(writer, sheet_name='All Jobs', index=False)
        format_worksheet(writer.sheets['All Jobs'], link_col_idx=5)
        
        # Write PG&E tab
        if not df_pge.empty:
            df_pge.to_excel(writer, sheet_name='PG&E', index=False)
            format_worksheet(writer.sheets['PG&E'], link_col_idx=5)
        
        # Write SMUD tab
        if not df_smud.empty:
            df_smud.to_excel(writer, sheet_name='SMUD', index=False)
            format_worksheet(writer.sheets['SMUD'], link_col_idx=5)
    
    print(f"\n✅ Merged {len(df_all)} jobs into {output_file}")
    print(f"   Tabs: All Jobs ({len(df_all)}), PG&E ({len(df_pge)}), SMUD ({len(df_smud)})")
    print(f"\n📝 Old files can be deleted if you want:")
    if os.path.exists("pge_jobs_master.xlsx"):
        print(f"   - pge_jobs_master.xlsx")
    
    return output_file


if __name__ == "__main__":
    merge_excel_files()
